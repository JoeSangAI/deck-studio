#!/usr/bin/env python3
"""Audit font families, inherited runs, and font-size tokens in a PPTX."""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from pptx import Presentation


def parse_csv(value: str | None) -> list[str]:
    if not value:
        return []
    return [item.strip() for item in value.split(",") if item.strip()]


def parse_sizes(value: str | None) -> list[float]:
    return [float(item) for item in parse_csv(value)]


def parse_slide_set(value: str | None) -> set[int]:
    slides: set[int] = set()
    for item in parse_csv(value):
        if "-" in item:
            start, end = (int(part.strip()) for part in item.split("-", 1))
            if start > end:
                raise ValueError(f"invalid slide range: {item}")
            slides.update(range(start, end + 1))
        else:
            slides.add(int(item))
    return slides


def near_allowed(size: float, allowed: list[float], tolerance: float) -> bool:
    return not allowed or any(abs(size - item) <= tolerance for item in allowed)


def audit(args: argparse.Namespace) -> tuple[dict, bool]:
    prs = Presentation(args.pptx)
    allowed_fonts = set(parse_csv(args.allowed_fonts))
    allowed_sizes = parse_sizes(args.allowed_sizes)
    excluded = parse_slide_set(args.exclude_slides)
    included = parse_slide_set(getattr(args, "include_slides", None))

    font_chars: Counter[str] = Counter()
    size_chars: Counter[str] = Counter()
    violations: list[dict] = []
    per_slide: dict[int, Counter[str]] = defaultdict(Counter)

    for slide_num, slide in enumerate(prs.slides, start=1):
        if (included and slide_num not in included) or slide_num in excluded:
            continue
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    text = run.text.strip()
                    if not text:
                        continue
                    char_count = len(text)
                    font_name = run.font.name or "(inherit)"
                    size = round(run.font.size.pt, 2) if run.font.size else None
                    size_key = "(inherit)" if size is None else f"{size:g}"
                    font_chars[font_name] += char_count
                    size_chars[size_key] += char_count
                    per_slide[slide_num][f"{font_name}|{size_key}"] += char_count

                    reasons: list[str] = []
                    if args.fail_on_inherited and (font_name == "(inherit)" or size is None):
                        reasons.append("inherited typography")
                    if allowed_fonts and font_name != "(inherit)" and font_name not in allowed_fonts:
                        reasons.append(f"font not allowed: {font_name}")
                    if size is not None and not near_allowed(size, allowed_sizes, args.tolerance):
                        reasons.append(f"size not allowed: {size:g}")
                    if reasons:
                        violations.append(
                            {
                                "slide": slide_num,
                                "shape": shape.name,
                                "text": text[:80],
                                "font": font_name,
                                "size": size,
                                "reasons": reasons,
                            }
                        )

    report = {
        "pptx": str(Path(args.pptx).resolve()),
        "slides": len(prs.slides),
        "included_slides": sorted(included),
        "excluded_slides": sorted(excluded),
        "font_chars": dict(font_chars.most_common()),
        "size_chars": dict(size_chars.most_common()),
        "per_slide_tokens": {
            str(slide): dict(tokens.most_common()) for slide, tokens in sorted(per_slide.items())
        },
        "violations_count": len(violations),
        "violations": violations,
    }
    return report, not violations


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", help="PowerPoint file to audit")
    parser.add_argument("--allowed-fonts", help="Comma-separated explicit font names")
    parser.add_argument("--allowed-sizes", help="Comma-separated point sizes")
    parser.add_argument("--tolerance", type=float, default=0.15, help="Point-size tolerance")
    parser.add_argument("--fail-on-inherited", action="store_true")
    parser.add_argument(
        "--include-slides",
        help="Only audit these 1-based slides; supports ranges such as 2,6,10-11",
    )
    parser.add_argument("--exclude-slides", help="Comma-separated 1-based slide numbers")
    parser.add_argument("--json-output", help="Optional JSON report path")
    args = parser.parse_args()

    report, ok = audit(args)
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.json_output:
        Path(args.json_output).write_text(payload + "\n", encoding="utf-8")
    print(payload)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
