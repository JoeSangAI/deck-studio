#!/usr/bin/env python3
"""
Generic PPTX audit for the deck-studio PPT workspace.

Use this before design work. It produces a lightweight JSON report with:
- deck size and slide count
- per-slide text/image/shape density
- title guess and page archetype guess
- font/color/media risk summaries
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation


PICTURE_TYPE = 13
NOISE_TEXT_PATTERNS = [
    re.compile(r"\boverall_\d+_\d+"),
    re.compile(r"\bcolumns_\d+_\d+"),
    re.compile(r"^[A-Za-z]+_\d+_\d+(?:\s+[A-Za-z]+_\d+_\d+)*$"),
]


def emu_to_inch(value: int) -> float:
    return value / 914400


def text_of(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    text = (shape.text_frame.text or "").strip()
    if is_noise_text(text):
        return ""
    return text


def is_noise_text(text: str) -> bool:
    compact = " ".join((text or "").split())
    if not compact:
        return True
    return any(pattern.search(compact) for pattern in NOISE_TEXT_PATTERNS)


def title_guess(slide) -> str:
    candidates = []
    for shape in slide.shapes:
        text = text_of(shape)
        if not text:
            continue
        w = emu_to_inch(shape.width)
        h = emu_to_inch(shape.height)
        top = emu_to_inch(shape.top)
        score = 0
        if top < 1.2:
            score += 4
        if w > 5:
            score += 3
        if len(text) <= 40:
            score += 2
        if h < 1.2:
            score += 1
        candidates.append((score, top, text.replace("\n", " | ")))
    if not candidates:
        return ""
    candidates.sort(key=lambda item: (-item[0], item[1]))
    return candidates[0][2]


def color_to_hex(color) -> str | None:
    try:
        rgb = color.rgb
        if rgb:
            return f"#{rgb}"
    except Exception:
        return None
    return None


def slide_stats(slide, index: int) -> dict:
    texts = [text_of(s) for s in slide.shapes if text_of(s)]
    text_chars = sum(len(t) for t in texts)
    pictures = [s for s in slide.shapes if s.shape_type == PICTURE_TYPE]
    text_shapes = [s for s in slide.shapes if getattr(s, "has_text_frame", False) and text_of(s)]
    shape_count = len(slide.shapes)
    title = title_guess(slide)
    archetype = classify_archetype(index, title, text_chars, len(pictures), shape_count, texts)
    return {
        "index": index,
        "title": title,
        "archetype_guess": archetype,
        "shape_count": shape_count,
        "text_shape_count": len(text_shapes),
        "picture_count": len(pictures),
        "text_chars": text_chars,
        "density": density_label(text_chars, shape_count, len(pictures)),
    }


def density_label(text_chars: int, shape_count: int, picture_count: int) -> str:
    if text_chars > 650 or shape_count > 45:
        return "dense"
    if picture_count >= 4:
        return "image-heavy"
    if text_chars < 80 and shape_count < 12:
        return "sparse"
    return "medium"


def classify_archetype(index: int, title: str, text_chars: int, picture_count: int, shape_count: int, texts: list[str]) -> str:
    joined = "\n".join(texts)
    title_norm = title.replace(" ", "")
    if index == 1:
        return "cover"
    if any(k in title_norm for k in ["谢谢", "感谢", "期待合作", "THANK"]):
        return "closing"
    if any(k in title_norm for k in ["目录", "总览", "指南", "框架"]):
        return "overview"
    if any(k in title_norm for k in ["方案", "策略", "合作"]):
        if any(k in joined for k in ["6月", "7月", "8月", "9月", "10月", "11月", "12月", "年度"]):
            return "timeline-plan"
        return "strategy-plan"
    if any(k in joined for k in ["月", "年度", "周"]) and shape_count > 20:
        return "timeline"
    if any(k in title_norm for k in ["案例", "发布会", "宣讲", "传播"]) or picture_count >= 3:
        return "case-image"
    if any(k in joined for k in ["费用", "预算", "金额", "花费", "覆盖"]):
        return "data-or-budget"
    if text_chars > 650:
        return "dense-text"
    if picture_count == 0 and text_chars < 160:
        return "section"
    return "mixed"


def collect_fonts_and_colors(prs: Presentation) -> tuple[Counter, Counter]:
    fonts = Counter()
    colors = Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        fonts[run.font.name] += 1
                    if run.font.size:
                        fonts[f"size:{run.font.size.pt:.1f}"] += 1
                    color = color_to_hex(run.font.color)
                    if color:
                        colors[color] += 1
    return fonts, colors


def media_risks(path: Path) -> list[str]:
    risks = []
    with ZipFile(path) as z:
        for name in z.namelist():
            if not name.startswith("ppt/media/"):
                continue
            suffix = Path(name).suffix.lower()
            head = z.read(name)[:256].lstrip().lower()
            if suffix in {".bin", ".svg"} or head.startswith(b"<svg") or b"<svg" in head[:128]:
                risks.append(name)
    return risks


def guess_scenario(slides: list[dict]) -> str:
    all_text = "\n".join(s["title"] for s in slides)
    if any(k in all_text for k in ["方案", "合作", "客户", "传播", "经销商"]):
        return "client-proposal"
    if any(k in all_text for k in ["经营", "复盘", "指标", "KPI", "预算"]):
        return "internal-report"
    if any(k in all_text for k in ["产品", "品牌", "发布", "上市"]):
        return "product-brand-launch"
    if any(k in all_text for k in ["投资", "财务", "市场规模", "增长"]):
        return "investor-data-report"
    if any(k in all_text for k in ["培训", "课程", "学习", "练习"]):
        return "training"
    return "general-business"


def audit(path: Path) -> dict:
    prs = Presentation(str(path))
    slides = [slide_stats(slide, i + 1) for i, slide in enumerate(prs.slides)]
    fonts, colors = collect_fonts_and_colors(prs)
    return {
        "path": str(path),
        "slide_count": len(prs.slides),
        "size_inches": {
            "width": round(emu_to_inch(prs.slide_width), 2),
            "height": round(emu_to_inch(prs.slide_height), 2),
        },
        "scenario_guess": guess_scenario(slides),
        "slides": slides,
        "top_fonts_and_sizes": fonts.most_common(12),
        "top_colors": colors.most_common(12),
        "media_risks": media_risks(path),
    }


def print_summary(report: dict) -> None:
    print(f"PPT AUDIT: {report['path']}")
    print(f"slides={report['slide_count']} size={report['size_inches']['width']}x{report['size_inches']['height']} scenario={report['scenario_guess']}")
    if report["media_risks"]:
        print(f"media risks: {', '.join(report['media_risks'])}")
    print("\nslides:")
    for s in report["slides"]:
        print(
            f"  P{s['index']:02d} {s['archetype_guess']:<14} {s['density']:<11} "
            f"text={s['text_chars']:<4} pics={s['picture_count']:<2} shapes={s['shape_count']:<3} title={s['title'][:60]}"
        )


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit a PPTX before polishing.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--out", type=Path, help="Write JSON report to this path")
    args = parser.parse_args()

    report = audit(args.pptx)
    print_summary(report)
    if args.out:
        args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nwrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
