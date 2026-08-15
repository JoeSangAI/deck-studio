"""
窄容器两行文本工具
- 适用场景：日期标签、狭窄时间槽等"主信息 + 副信息"组合
- 排版：上行（主信息，10pt）+ 下行（副信息，8pt），居中对齐
"""
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN


def set_two_line_text(text_frame, main_text, sub_text,
                       main_size=10, sub_size=8,
                       main_bold=False, sub_bold=False,
                       color=None):
    """
    设置文本框为两行排版

    text_frame: python-pptx 的 text_frame
    main_text: 主信息（如日期 "12.21-12.27"）
    sub_text: 副信息（如周数 "(1周)"）
    main_size: 主信息字号（pt）
    sub_size: 副信息字号（pt）
    main_bold: 主信息是否加粗
    sub_bold: 副信息是否加粗
    color: 文字颜色（RGB 对象）
    """
    # 清空现有内容
    text_frame.clear()

    # 第一行：主信息
    p1 = text_frame.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = main_text
    r1.font.size = Pt(main_size)
    r1.font.bold = main_bold
    if color is not None:
        r1.font.color.rgb = color

    # 第二行：副信息
    p2 = text_frame.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = sub_text
    r2.font.size = Pt(sub_size)
    r2.font.bold = sub_bold
    if color is not None:
        r2.font.color.rgb = color


def date_with_weeks(years_pair, weeks_count):
    """
    快捷方法：把 (起始日期, 结束日期) + 周数 转成两行文本格式

    years_pair: ("12.21", "12.27") 或 ("6.1", "6.30")
    weeks_count: 1, 2, 3, 4 等
    return: (main_text, sub_text)
    """
    start, end = years_pair
    main_text = f"{start}-{end}"
    sub_text = f"({weeks_count}周)"
    return main_text, sub_text


# 字号预设（按容器宽度选）
SIZE_PRESETS = {
    "narrow": {"main": 9, "sub": 7},    # 容器 < 0.7"
    "medium": {"main": 10, "sub": 8},   # 容器 0.7"-1.2"
    "wide": {"main": 12, "sub": 9},     # 容器 > 1.2"
}


def get_preset(preset_name="medium"):
    return SIZE_PRESETS.get(preset_name, SIZE_PRESETS["medium"])


if __name__ == "__main__":
    # 示例
    main, sub = date_with_weeks(("12.21", "12.27"), 1)
    print(f"主信息: '{main}'")
    print(f"副信息: '{sub}'")

    print("\n字号预设:")
    for name, sizes in SIZE_PRESETS.items():
        print(f"  {name}: 主 {sizes['main']}pt / 副 {sizes['sub']}pt")
