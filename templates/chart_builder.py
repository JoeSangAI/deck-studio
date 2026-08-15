#!/usr/bin/env python3
"""
Insert editable PowerPoint charts from a small JSON specification.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt


class ChartSpecError(ValueError):
    """Raised when chart_spec.json is invalid."""


CHART_TYPES = {
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "pie": XL_CHART_TYPE.PIE,
    "radar": XL_CHART_TYPE.RADAR_MARKERS,
}


DEFAULT_COLORS = ["#C9A84C", "#4A3F35", "#4E79A7", "#59A14F", "#E15759", "#B07AA1"]


def _hex_to_rgb(value: str) -> RGBColor:
    raw = value.strip().lstrip("#")
    if len(raw) != 6:
        raise ChartSpecError(f"invalid color {value!r}; expected #RRGGBB")
    try:
        return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))
    except ValueError as exc:
        raise ChartSpecError(f"invalid color {value!r}; expected #RRGGBB") from exc


def _palette(spec: dict[str, Any], primary: str | None, secondary: str | None, accent: str | None) -> list[RGBColor]:
    values = []
    if primary:
        values.append(primary)
    if secondary:
        values.append(secondary)
    if accent:
        values.append(accent)
    values.extend(spec.get("palette", []))
    values.extend(DEFAULT_COLORS)

    seen = set()
    out = []
    for value in values:
        if not value or value in seen:
            continue
        seen.add(value)
        out.append(_hex_to_rgb(value))
    return out


def _require_number(chart: dict[str, Any], field: str) -> float:
    value = chart.get(field)
    if not isinstance(value, (int, float)):
        raise ChartSpecError(f"{field} must be a number of inches")
    if value < 0:
        raise ChartSpecError(f"{field} must be non-negative")
    return float(value)


def _validate_chart(chart: dict[str, Any], slide_count: int) -> None:
    if not isinstance(chart, dict):
        raise ChartSpecError("each chart must be an object")

    slide = chart.get("slide")
    if not isinstance(slide, int) or slide < 1 or slide > slide_count:
        raise ChartSpecError(f"slide must be between 1 and {slide_count}")

    kind = chart.get("kind")
    if kind not in CHART_TYPES:
        raise ChartSpecError(f"kind must be one of: {', '.join(CHART_TYPES)}")

    for field in ["left", "top", "width", "height"]:
        _require_number(chart, field)
    if chart["width"] <= 0 or chart["height"] <= 0:
        raise ChartSpecError("width and height must be greater than zero")

    categories = chart.get("categories")
    if not isinstance(categories, list) or not categories:
        raise ChartSpecError("categories must be a non-empty list")
    if not all(isinstance(item, (str, int, float)) for item in categories):
        raise ChartSpecError("categories must contain only text or numbers")

    series = chart.get("series")
    if not isinstance(series, list) or not series:
        raise ChartSpecError("series must be a non-empty list")

    for item in series:
        if not isinstance(item, dict):
            raise ChartSpecError("each series must be an object")
        name = item.get("name")
        values = item.get("values")
        if not isinstance(name, str) or not name.strip():
            raise ChartSpecError("each series must have a non-empty name")
        if not isinstance(values, list) or len(values) != len(categories):
            raise ChartSpecError("series values must be the same length as categories")
        if not all(isinstance(value, (int, float)) for value in values):
            raise ChartSpecError("series values must contain only numbers")

    if kind == "pie" and len(series) != 1:
        raise ChartSpecError("pie charts require exactly one series")


def _validate_spec(spec: dict[str, Any], slide_count: int) -> list[dict[str, Any]]:
    if not isinstance(spec, dict):
        raise ChartSpecError("spec root must be an object")
    charts = spec.get("charts")
    if not isinstance(charts, list) or not charts:
        raise ChartSpecError("spec must contain a non-empty charts list")
    for chart in charts:
        _validate_chart(chart, slide_count)
    return charts


def _chart_data(chart: dict[str, Any]) -> CategoryChartData:
    data = CategoryChartData()
    data.categories = [str(item) for item in chart["categories"]]
    for series in chart["series"]:
        data.add_series(series["name"], tuple(series["values"]))
    return data


def _style_chart(chart_obj, palette: list[RGBColor], has_multiple_series: bool) -> None:
    chart_obj.has_legend = has_multiple_series
    if has_multiple_series:
        chart_obj.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart_obj.legend.include_in_layout = False

    chart_obj.font.size = Pt(9)
    chart_obj.font.name = "Arial"
    try:
        chart_obj.category_axis.tick_labels.font.size = Pt(8)
        chart_obj.value_axis.tick_labels.font.size = Pt(8)
        chart_obj.value_axis.has_major_gridlines = True
    except Exception:
        pass

    for index, series in enumerate(chart_obj.series):
        color = palette[index % len(palette)]
        try:
            series.format.line.color.rgb = color
            series.format.line.width = Pt(1.8)
        except Exception:
            pass
        try:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
        except Exception:
            pass


def _add_chart(slide, chart: dict[str, Any], palette: list[RGBColor]) -> None:
    graphic_frame = slide.shapes.add_chart(
        CHART_TYPES[chart["kind"]],
        Inches(float(chart["left"])),
        Inches(float(chart["top"])),
        Inches(float(chart["width"])),
        Inches(float(chart["height"])),
        _chart_data(chart),
    )
    chart_obj = graphic_frame.chart
    if chart.get("title"):
        chart_obj.has_title = True
        chart_obj.chart_title.text_frame.text = str(chart["title"])
        chart_obj.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    _style_chart(chart_obj, palette, has_multiple_series=len(chart["series"]) > 1)


def apply_chart_spec(
    input_path: str | Path,
    output_path: str | Path,
    spec: dict[str, Any],
    *,
    primary: str | None = "#C9A84C",
    secondary: str | None = "#4A3F35",
    accent: str | None = None,
) -> None:
    source = Path(input_path)
    output = Path(output_path)
    prs = Presentation(str(source))
    charts = _validate_spec(spec, len(prs.slides))
    palette = _palette(spec, primary, secondary, accent)

    for chart in charts:
        slide = prs.slides[int(chart["slide"]) - 1]
        _add_chart(slide, chart, palette)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Insert editable charts into a PPTX from chart_spec.json.")
    parser.add_argument("input", type=Path, help="Input PPTX")
    parser.add_argument("--output", "-o", required=True, type=Path, help="Output PPTX")
    parser.add_argument("--spec", required=True, type=Path, help="Chart spec JSON")
    parser.add_argument("--primary", default="#C9A84C", help="Primary color #RRGGBB")
    parser.add_argument("--secondary", default="#4A3F35", help="Secondary color #RRGGBB")
    parser.add_argument("--accent", default=None, help="Accent color #RRGGBB")
    args = parser.parse_args(argv)

    try:
        spec = json.loads(args.spec.read_text(encoding="utf-8"))
        apply_chart_spec(
            args.input,
            args.output,
            spec,
            primary=args.primary,
            secondary=args.secondary,
            accent=args.accent,
        )
    except (OSError, json.JSONDecodeError, ChartSpecError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    print(f"wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
