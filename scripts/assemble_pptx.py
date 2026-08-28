# -*- coding: utf-8 -*-
"""Assemble a legacy image-only draft PPTX from deck.json.

Usage: python assemble_pptx.py deck.json --out draft.pptx
Formal and mixed-type exports belong to PPT God's Workflow Core. This helper
preflights every image before creating an output; it never inserts blank pages.
Focus Media slides are rejected unless their Output Check report matches the
exact current PNG and reference assets.
"""
import argparse
import json
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

SKILL_ROOT = Path(__file__).resolve().parents[1]
if str(SKILL_ROOT) not in sys.path:
    sys.path.insert(0, str(SKILL_ROOT))

from templates.verify_focusmedia_output import validate_media_output

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

SLIDE_W, SLIDE_H = 12192000, 6858000  # 16:9 EMU


def main():
    parser = argparse.ArgumentParser(description="assemble full-bleed PPTX from generated PNGs")
    parser.add_argument("deck", help="path to deck.json")
    parser.add_argument("--out", required=True, help="output .pptx path")
    args = parser.parse_args()

    cfg = json.load(open(args.deck, encoding="utf-8"))
    base_dir = Path(args.deck).resolve().parent
    outdir = base_dir / cfg.get("outdir", "png")
    width = int(cfg.get("size", "3840x2160").lower().split("x")[0])
    min_bytes = 300000 if width >= 3000 else 40000

    media_errors = []
    for slide_def in cfg["slides"]:
        if slide_def.get("specialist_route") != "focusmedia-image-gen":
            continue
        png = outdir / (slide_def["id"] + ".png")
        media_errors.extend(
            validate_media_output(
                slide_def,
                base_dir=base_dir,
                output_path=png,
            )
        )
    if media_errors:
        for error in media_errors:
            print("MEDIA VALIDATION ERROR: " + error)
        sys.exit(1)

    unsupported = [
        slide_def["id"]
        for slide_def in cfg["slides"]
        if slide_def.get("production_type") in {"hybrid", "native_editable"}
    ]
    if unsupported:
        print(
            "PPT GOD ASSEMBLY REQUIRED: hybrid/native_editable slides: "
            + ", ".join(unsupported)
        )
        sys.exit(1)

    missing = []
    for slide_def in cfg["slides"]:
        png = outdir / (slide_def["id"] + ".png")
        if not png.exists() or png.stat().st_size <= min_bytes:
            missing.append(slide_def["id"])
    if missing:
        print("MISSING: " + ", ".join(missing))
        sys.exit(1)

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    blank = prs.slide_layouts[6]

    placed = []
    for slide_def in cfg["slides"]:
        sid = slide_def["id"]
        png = outdir / (sid + ".png")
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        placed.append(sid)

    prs.save(args.out)
    print("saved: %s  (%.1f MB)" % (args.out, os.path.getsize(args.out) / 1e6))
    print("placed %d / %d" % (len(placed), len(cfg["slides"])))


if __name__ == "__main__":
    main()
