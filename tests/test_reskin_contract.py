import importlib.util
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


MODULE_PATH = Path(__file__).parents[1] / "templates" / "reskin_cli.py"
SPEC = importlib.util.spec_from_file_location("reskin_cli", MODULE_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader
SPEC.loader.exec_module(MODULE)


def make_deck(path: Path, font_name: str = "Original Font") -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(
        Inches(0.5),
        Inches(0.3),
        Inches(11.2),
        Inches(0.6),
    )
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "这是页面主标题"
    run.font.name = font_name
    run.font.size = Pt(24)
    prs.save(path)


def title_font(path: Path) -> str | None:
    prs = Presentation(path)
    return prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name


def test_reskin_has_no_project_specific_default_skips(tmp_path):
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    make_deck(source)

    result = MODULE.main([
        str(source),
        "--output",
        str(output),
        "--title-font",
        "New Font",
    ])

    assert result == 0
    assert title_font(output) == "New Font"


def test_explicitly_skipped_slide_is_not_restyled(tmp_path):
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    make_deck(source)

    result = MODULE.main([
        str(source),
        "--output",
        str(output),
        "--title-font",
        "New Font",
        "--skip-slides",
        "1",
    ])

    assert result == 0
    assert title_font(output) == "Original Font"


def test_processing_error_is_nonzero_and_writes_no_output(
    monkeypatch,
    tmp_path,
):
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    make_deck(source)
    monkeypatch.setattr(
        MODULE,
        "process_slide",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("boom")),
    )

    result = MODULE.main([
        str(source),
        "--output",
        str(output),
    ])

    assert result == 1
    assert not output.exists()


def test_preservation_failure_returns_nonzero(monkeypatch, tmp_path):
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    make_deck(source)
    monkeypatch.setattr(
        MODULE,
        "verify_preservation",
        lambda *args: ["Slide 1: text lost"],
    )

    result = MODULE.main([
        str(source),
        "--output",
        str(output),
    ])

    assert result == 1
    assert output.exists()


def test_ai_background_flag_fails_instead_of_claiming_unused_output(tmp_path):
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    make_deck(source)

    result = MODULE.main([
        str(source),
        "--output",
        str(output),
        "--ai-bg",
    ])

    assert result == 2
    assert not output.exists()
