from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from templates.preserve_ratio import (
    apply_to_pic,
    fit_contain,
    unify_height,
    unify_width,
)
from templates.two_line_text import date_with_weeks, get_preset, set_two_line_text


def test_ratio_helpers_preserve_dimensions():
    assert unify_height((1920, 1080), 4.0) == pytest.approx((64 / 9, 4.0))
    assert unify_width((1920, 1080), 8.0) == pytest.approx((8.0, 4.5))
    assert fit_contain((1920, 1080), 5.0, 4.0) == pytest.approx((5.0, 2.8125))


def test_apply_to_pic_sets_powerpoint_dimensions():
    picture = SimpleNamespace(width=0, height=0)
    apply_to_pic(picture, 5.0, 2.5)
    assert picture.width == Inches(5.0)
    assert picture.height == Inches(2.5)


def test_two_line_text_builds_editable_paragraphs():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_frame = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame

    set_two_line_text(
        text_frame,
        "12.21-12.27",
        "(1周)",
        main_size=11,
        sub_size=8,
        main_bold=True,
        color=RGBColor(10, 20, 30),
    )

    assert len(text_frame.paragraphs) == 2
    assert [paragraph.text for paragraph in text_frame.paragraphs] == [
        "12.21-12.27",
        "(1周)",
    ]
    assert all(paragraph.alignment == PP_ALIGN.CENTER for paragraph in text_frame.paragraphs)
    assert text_frame.paragraphs[0].runs[0].font.size == Pt(11)
    assert text_frame.paragraphs[0].runs[0].font.bold is True
    assert text_frame.paragraphs[1].runs[0].font.size == Pt(8)
    assert text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(10, 20, 30)


def test_two_line_text_helpers_are_stable():
    assert date_with_weeks(("6.1", "6.30"), 4) == ("6.1-6.30", "(4周)")
    assert get_preset("narrow") == {"main": 9, "sub": 7}
    assert get_preset("unknown") == get_preset("medium")
