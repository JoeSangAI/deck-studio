import argparse
import importlib.util
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


MODULE_PATH = Path(__file__).parents[1] / "templates" / "typography_audit.py"
SPEC = importlib.util.spec_from_file_location("typography_audit", MODULE_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader
SPEC.loader.exec_module(MODULE)


def make_deck(path: Path, font: str | None, size: float | None) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 1000000, 500000)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "测试 Typography"
    if font:
        run.font.name = font
    if size:
        run.font.size = Pt(size)
    prs.save(path)


def args_for(path: Path, **overrides):
    values = {
        "pptx": str(path),
        "allowed_fonts": "PingFang SC",
        "allowed_sizes": "12,18,30",
        "tolerance": 0.15,
        "fail_on_inherited": True,
        "include_slides": None,
        "exclude_slides": None,
        "json_output": None,
    }
    values.update(overrides)
    return argparse.Namespace(**values)


def test_typography_audit_accepts_allowed_tokens(tmp_path):
    path = tmp_path / "ok.pptx"
    make_deck(path, "PingFang SC", 18)
    report, ok = MODULE.audit(args_for(path))
    assert ok
    assert report["violations_count"] == 0


def test_typography_audit_rejects_inherited_and_fragment_size(tmp_path):
    path = tmp_path / "bad.pptx"
    make_deck(path, None, 17.2)
    report, ok = MODULE.audit(args_for(path))
    assert not ok
    assert report["violations_count"] == 1
    reasons = report["violations"][0]["reasons"]
    assert "inherited typography" in reasons
    assert "size not allowed: 17.2" in reasons


def test_parse_slide_set_supports_ranges():
    assert MODULE.parse_slide_set("2,6,10-12") == {2, 6, 10, 11, 12}


def test_typography_audit_can_gate_only_rewritten_slides(tmp_path):
    path = tmp_path / "scoped.pptx"
    prs = Presentation()

    old_slide = prs.slides.add_slide(prs.slide_layouts[6])
    old_box = old_slide.shapes.add_textbox(0, 0, 1000000, 500000)
    old_run = old_box.text_frame.paragraphs[0].add_run()
    old_run.text = "保留页旧字体"
    old_run.font.name = "SimSun"
    old_run.font.size = Pt(18)

    new_slide = prs.slides.add_slide(prs.slide_layouts[6])
    new_box = new_slide.shapes.add_textbox(0, 0, 1000000, 500000)
    new_run = new_box.text_frame.paragraphs[0].add_run()
    new_run.text = "重做页新字体"
    new_run.font.name = "微软雅黑"
    new_run.font.size = Pt(18)
    prs.save(path)

    report, ok = MODULE.audit(
        args_for(
            path,
            allowed_fonts="微软雅黑",
            include_slides="2",
        )
    )
    assert ok
    assert report["included_slides"] == [2]
    assert report["violations_count"] == 0

    report, ok = MODULE.audit(
        args_for(
            path,
            allowed_fonts="微软雅黑",
            include_slides="1",
        )
    )
    assert not ok
    assert report["violations"][0]["reasons"] == ["font not allowed: SimSun"]
