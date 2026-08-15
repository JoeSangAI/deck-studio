from pathlib import Path
import json
import sys
from zipfile import ZipFile

import pytest
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "templates"))

from chart_builder import ChartSpecError, apply_chart_spec


def make_deck(path: Path, slide_count: int = 2) -> Path:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    while len(prs.slides) < slide_count:
        prs.slides.add_slide(blank)
    prs.save(path)
    return path


def read_chart_xml(path: Path) -> str:
    with ZipFile(path) as archive:
        chart_names = sorted(
            name for name in archive.namelist()
            if name.startswith("ppt/charts/chart") and name.endswith(".xml")
        )
        assert chart_names
        return "\n".join(archive.read(name).decode("utf-8") for name in chart_names)


def test_line_chart_is_inserted_as_editable_ppt_chart(tmp_path):
    source = make_deck(tmp_path / "source.pptx")
    output = tmp_path / "output.pptx"
    spec = {
        "charts": [{
            "slide": 1,
            "kind": "line",
            "left": 1.0,
            "top": 1.0,
            "width": 5.0,
            "height": 3.0,
            "categories": ["Q1", "Q2", "Q3"],
            "series": [{"name": "Revenue", "values": [10, 15, 21]}],
        }]
    }

    apply_chart_spec(source, output, spec)

    chart_xml = read_chart_xml(output)
    assert "<c:lineChart>" in chart_xml
    assert "Revenue" in chart_xml


def test_radar_chart_is_inserted_as_editable_ppt_chart(tmp_path):
    source = make_deck(tmp_path / "source.pptx")
    output = tmp_path / "output.pptx"
    spec = {
        "charts": [{
            "slide": 1,
            "kind": "radar",
            "left": 1.0,
            "top": 1.0,
            "width": 5.0,
            "height": 3.5,
            "categories": ["Coverage", "Cost", "Speed"],
            "series": [
                {"name": "Plan A", "values": [80, 70, 90]},
                {"name": "Plan B", "values": [72, 82, 75]},
            ],
        }]
    }

    apply_chart_spec(source, output, spec)

    chart_xml = read_chart_xml(output)
    assert "<c:radarChart>" in chart_xml
    assert "Plan A" in chart_xml


def test_series_length_mismatch_raises_and_does_not_write_output(tmp_path):
    source = make_deck(tmp_path / "source.pptx")
    output = tmp_path / "output.pptx"
    spec = {
        "charts": [{
            "slide": 1,
            "kind": "line",
            "left": 1.0,
            "top": 1.0,
            "width": 5.0,
            "height": 3.0,
            "categories": ["Q1", "Q2", "Q3"],
            "series": [{"name": "Revenue", "values": [10, 15]}],
        }]
    }

    with pytest.raises(ChartSpecError, match="same length"):
        apply_chart_spec(source, output, spec)

    assert not output.exists()


def test_slide_index_out_of_range_raises_and_does_not_write_output(tmp_path):
    source = make_deck(tmp_path / "source.pptx", slide_count=1)
    output = tmp_path / "output.pptx"
    spec = {
        "charts": [{
            "slide": 3,
            "kind": "bar",
            "left": 1.0,
            "top": 1.0,
            "width": 5.0,
            "height": 3.0,
            "categories": ["A", "B"],
            "series": [{"name": "Score", "values": [1, 2]}],
        }]
    }

    with pytest.raises(ChartSpecError, match="slide"):
        apply_chart_spec(source, output, spec)

    assert not output.exists()


def test_cli_reads_json_spec(tmp_path):
    source = make_deck(tmp_path / "source.pptx")
    output = tmp_path / "output.pptx"
    spec_path = tmp_path / "chart_spec.json"
    spec_path.write_text(json.dumps({
        "charts": [{
            "slide": 1,
            "kind": "column",
            "left": 1.0,
            "top": 1.0,
            "width": 5.0,
            "height": 3.0,
            "categories": ["A", "B"],
            "series": [{"name": "Score", "values": [1, 2]}],
        }]
    }), encoding="utf-8")

    from chart_builder import main

    assert main([str(source), "--output", str(output), "--spec", str(spec_path)]) == 0
    assert "<c:barChart>" in read_chart_xml(output)
